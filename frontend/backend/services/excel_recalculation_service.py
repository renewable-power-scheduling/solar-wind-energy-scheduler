import logging
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import load_workbook

LOGGER = logging.getLogger(__name__)

REQUIRED_FINAL_CHARGE_CELLS = {
    "Final_Summary": ("D7", "E7", "F7"),
}


def _find_libreoffice_binary() -> str:
    for candidate in ("soffice", "libreoffice"):
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    raise ValueError("LibreOffice headless executable not found. Install soffice or libreoffice.")


def _validate_recalculated_workbook(
    workbook_path: Path,
    required_final_charge_cells: Mapping[str, Sequence[str]] | None = None,
    allow_formula_fallback: bool = False,
) -> None:
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    formula_wb = load_workbook(workbook_path, data_only=False, read_only=False) if allow_formula_fallback else None
    required_cells = REQUIRED_FINAL_CHARGE_CELLS if required_final_charge_cells is None else required_final_charge_cells
    missing = []
    for sheet_name, cell_refs in required_cells.items():
        if sheet_name not in wb.sheetnames:
            missing.append(sheet_name)
            continue
        ws = wb[sheet_name]
        formula_ws = formula_wb[sheet_name] if formula_wb is not None and sheet_name in formula_wb.sheetnames else None
        for cell_ref in cell_refs:
            value = ws[cell_ref].value
            if value in (None, ""):
                formula_value = formula_ws[cell_ref].value if formula_ws is not None else None
                if isinstance(formula_value, str) and formula_value.startswith("="):
                    continue
                missing.append(f"{sheet_name}!{cell_ref}")
    if missing:
        raise ValueError(
            "DSM Verification workbook was generated but Final Charges could not be recalculated: "
            + ", ".join(sorted(set(missing)))
        )


def recalculate_excel_workbook(
    input_path,
    output_path,
    timeout_seconds: int = 300,
    required_final_charge_cells: Mapping[str, Sequence[str]] | None = None,
    allow_formula_fallback: bool = False,
) -> Path:
    input_path = Path(input_path)
    output_path = Path(output_path)

    if not input_path.exists():
        raise FileNotFoundError(f"Workbook not found: {input_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    soffice = _find_libreoffice_binary()

    with tempfile.TemporaryDirectory(prefix="dsm_excel_recalc_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        staged_input = temp_dir / input_path.name
        temp_output_dir = temp_dir / "out"
        profile_dir = temp_dir / "profile"
        temp_output_dir.mkdir(parents=True, exist_ok=True)
        profile_dir.mkdir(parents=True, exist_ok=True)

        shutil.copy2(input_path, staged_input)
        profile_uri = profile_dir.resolve().as_uri()
        cmd = [
            soffice,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--nofirststartwizard",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(temp_output_dir),
            str(staged_input),
        ]
        LOGGER.info("DSM Verification: running spreadsheet recalculation via LibreOffice")
        completed = subprocess.run(
            cmd,
            check=False,
            capture_output=True,
            text=True,
            timeout=timeout_seconds,
        )
        if completed.returncode != 0:
            raise ValueError(
                "Excel recalculation failed: "
                f"exit_code={completed.returncode}; stdout={completed.stdout.strip()}; stderr={completed.stderr.strip()}"
            )

        produced = temp_output_dir / staged_input.name
        if not produced.exists():
            candidates = list(temp_output_dir.glob("*.xlsx"))
            if len(candidates) == 1:
                produced = candidates[0]
            else:
                raise ValueError(
                    "Excel recalculation completed but the recalculated workbook was not produced. "
                    f"Expected file: {produced}"
                )

        shutil.copy2(produced, output_path)

    if not output_path.exists():
        raise ValueError(f"Excel recalculation failed to create output workbook: {output_path}")

    _validate_recalculated_workbook(output_path, required_final_charge_cells, allow_formula_fallback=allow_formula_fallback)
    LOGGER.info("DSM Verification: recalculated workbook validated successfully")
    return output_path
