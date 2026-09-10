import os
import sys
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

BACKEND_DIR = os.path.dirname(os.path.dirname(__file__))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from database import Base
from models import DsmVerificationRun
from services.dsm_verification_service import (
    _populate_workbook,
    _validate_rewa_formula_cells,
    create_run,
    generate_run_workbook,
    get_active_template_row,
    parse_meter_uploads_for_generators,
    store_template,
    validate_run_inputs,
)
from services.excel_recalculation_service import recalculate_excel_workbook


def build_rewa_template_bytes(with_formulas: bool = False) -> bytes:
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    for sheet_name in ("Aggregated_Calculations", "Summary_WA", "Summary_WoA", "Final_Summary"):
        wb.create_sheet(sheet_name)

    if with_formulas:
        ws = wb["Aggregated_Calculations"]
        ws["G4"] = "=D4-C4"
        ws["AK4"] = "=AH4-AG4"
        ws["BR4"] = "=IF(BP4=0,0,(BJ4/BN4)*BP4)"
        ws["CG4"] = "=IF(BZ4=0,0,(BZ4/CB4)*CF4)"
        ws = wb["Summary_WA"]
        ws["D8"] = "=SUM(A1:A2)"
        ws["D10"] = "=D8-D9"
        ws["D20"] = "=D18-D19"
        ws = wb["Summary_WoA"]
        ws["D8"] = "=SUM(A1:A2)"
        ws["D13"] = "=SUM(A3:A4)"
        ws["D15"] = "=D13-D14"
        ws = wb["Final_Summary"]
        ws["D7"] = "=1+1"
        ws["D8"] = '=IF(D7>0,"Receivable","Payable")'
        ws["E7"] = "=2+2"
        ws["E8"] = '=IF(E7>0,"Receivable","Payable")'
        ws["F7"] = "=3+3"
        ws["F8"] = '=IF(F7>0,"Receivable","Payable")'

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_recalculated_workbook(path: Path, final_values=(101.25, 202.5, 303.75)) -> None:
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    wb.create_sheet("Final_Summary")
    ws = wb["Final_Summary"]
    ws["D7"], ws["E7"], ws["F7"] = final_values
    wb.save(path)


class DsmVerificationRecalculationTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)
        self.db = self.Session()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_rewa_formula_cells_are_present_before_recalculation(self):
        _validate_rewa_formula_cells(build_rewa_template_bytes(with_formulas=True))
        with self.assertRaises(ValueError):
            _validate_rewa_formula_cells(build_rewa_template_bytes(with_formulas=False))

    def test_templates_are_active_per_regulation(self):
        template_2014 = store_template(
            self.db,
            pss_code="REWA",
            filename="RUMS-DSM-Report(2014 Reg)-All-three.xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=build_rewa_template_bytes(with_formulas=False),
            regulation="2014",
        )
        template_2026 = store_template(
            self.db,
            pss_code="REWA",
            filename="Rums(2026).xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            content=build_rewa_template_bytes(with_formulas=False),
            regulation="2026",
        )

        self.assertEqual(get_active_template_row(self.db, "REWA", "2014").id, template_2014.id)
        self.assertEqual(get_active_template_row(self.db, "REWA", "2026").id, template_2026.id)
        self.assertTrue(template_2014.is_active)
        self.assertTrue(template_2026.is_active)

    def test_run_inputs_are_shared_across_regulations(self):
        with patch("services.dsm_verification_service.get_pss_config", return_value={"pss_name": "Rewa"}):
            run_2014 = create_run(
                self.db,
                pss_code="REWA",
                from_date=date(2026, 8, 10),
                to_date=date(2026, 8, 16),
                regulation="2014",
            )
            run_2026 = create_run(
                self.db,
                pss_code="REWA",
                from_date=date(2026, 8, 10),
                to_date=date(2026, 8, 16),
                regulation="2026",
            )

        self.assertEqual(run_2014.id, run_2026.id)
        self.assertEqual(run_2026.regulation, "2026")

    def test_recalculate_excel_workbook_populates_final_charges(self):
        with tempfile.TemporaryDirectory() as temp_dir_name:
            temp_dir = Path(temp_dir_name)
            input_path = temp_dir / "input.xlsx"
            output_path = temp_dir / "output.xlsx"
            build_recalculated_workbook(input_path)

            def fake_run(cmd, check, capture_output, text, timeout):
                out_dir = Path(cmd[cmd.index("--outdir") + 1])
                source_path = Path(cmd[-1])
                produced = out_dir / source_path.name
                build_recalculated_workbook(produced, (11.0, 22.0, 33.0))
                return SimpleNamespace(returncode=0, stdout="ok", stderr="")

            with patch("services.excel_recalculation_service.shutil.which", return_value="soffice"), patch(
                "services.excel_recalculation_service.subprocess.run",
                side_effect=fake_run,
            ):
                result = recalculate_excel_workbook(input_path, output_path)

            self.assertEqual(result, output_path)
            wb = load_workbook(output_path, data_only=True)
            ws = wb["Final_Summary"]
            self.assertEqual(ws["D7"].value, 11.0)
            self.assertEqual(ws["E7"].value, 22.0)
            self.assertEqual(ws["F7"].value, 33.0)

    def test_recalculate_excel_workbook_allows_rewa_formula_cells_without_cached_values(self):
        with tempfile.TemporaryDirectory() as temp_dir_name:
            temp_dir = Path(temp_dir_name)
            input_path = temp_dir / "input.xlsx"
            output_path = temp_dir / "output.xlsx"
            input_path.write_bytes(build_rewa_template_bytes(with_formulas=True))

            def fake_run(cmd, check, capture_output, text, timeout):
                out_dir = Path(cmd[cmd.index("--outdir") + 1])
                source_path = Path(cmd[-1])
                produced = out_dir / source_path.name
                produced.write_bytes(source_path.read_bytes())
                return SimpleNamespace(returncode=0, stdout="ok", stderr="")

            with patch("services.excel_recalculation_service.shutil.which", return_value="soffice"), patch(
                "services.excel_recalculation_service.subprocess.run",
                side_effect=fake_run,
            ):
                result = recalculate_excel_workbook(
                    input_path,
                    output_path,
                    required_final_charge_cells={"Final_Summary": ("D7", "E7", "F7")},
                    allow_formula_fallback=True,
                )

            self.assertEqual(result, output_path)
            wb = load_workbook(output_path, data_only=False)
            ws = wb["Final_Summary"]
            self.assertTrue(str(ws["D7"].value).startswith("="))
            self.assertTrue(str(ws["E7"].value).startswith("="))
            self.assertTrue(str(ws["F7"].value).startswith("="))

    def test_generate_run_workbook_recalculates_and_stores_final_output(self):
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            from_date=date(2026, 8, 3),
            to_date=date(2026, 8, 9),
            revision_number=1,
            status="DRAFT",
            meter_count_expected=7,
            meter_count_uploaded=7,
            schedule_sprng_count_uploaded=7,
            schedule_seit_count_uploaded=7,
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            validation_status="PENDING",
            created_by="tester",
        )
        self.db.add(run)
        self.db.commit()
        self.db.refresh(run)

        template_bytes = build_rewa_template_bytes(with_formulas=True)

        def fake_recalc(input_path, output_path, **_kwargs):
            build_recalculated_workbook(Path(output_path), (101.25, 202.5, 303.75))
            return Path(output_path)

        with patch("services.dsm_verification_service.get_pss_config", return_value={"workbook": {"sheet_name": "Aggregated_Calculations", "start_row": 4, "rows_per_day": 96, "days": 7, "input_columns": {"date": "A", "time_block": "B", "SPRNG": {"schedule": "C", "actual": "D", "avc": "E", "ppa": "F"}, "SEIT": {"schedule": "R", "actual": "S", "avc": "T", "ppa": "U"}}}}), patch(
            "services.dsm_verification_service.get_active_template_row",
            return_value=SimpleNamespace(id=1, version=1, template_binary=template_bytes),
        ), patch(
            "services.dsm_verification_service.validate_run_inputs",
            return_value={"ok": True, "missing": []},
        ), patch(
            "services.dsm_verification_service._parse_input_files_for_run",
            return_value=(
                {
                    "SPRNG": {date(2026, 8, 3): {1: 1.0, 2: 1.0}},
                    "SEIT": {date(2026, 8, 3): {1: 1.0, 2: 1.0}},
                },
                {
                    "SPRNG": {date(2026, 8, 3): {1: 1.0, 2: 1.0}},
                    "SEIT": {date(2026, 8, 3): {1: 1.0, 2: 1.0}},
                },
                {},
                [],
            ),
        ), patch(
            "services.dsm_verification_service.recalculate_excel_workbook",
            side_effect=fake_recalc,
        ) as recalc_mock:
            result = generate_run_workbook(self.db, run.id)

        self.assertEqual(recalc_mock.call_count, 1)
        self.assertEqual(result.status, "COMPLETED")
        self.assertEqual(result.generated_filename, "REWA_DSM_Verification(2014)_2026-08-03_to_2026-08-09.xlsx")
        self.assertIsNotNone(result.generated_binary)
        wb = load_workbook(BytesIO(result.generated_binary), data_only=True)
        ws = wb["Final_Summary"]
        self.assertEqual(ws["D7"].value, 101.25)
        self.assertEqual(ws["E7"].value, 202.5)
        self.assertEqual(ws["F7"].value, 303.75)

    def test_generate_run_workbook_uses_extended_rewa_recalculation_cells(self):
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            regulation="2026",
            from_date=date(2026, 8, 3),
            to_date=date(2026, 8, 9),
            revision_number=1,
            status="DRAFT",
            meter_count_expected=7,
            meter_count_uploaded=7,
            schedule_sprng_count_uploaded=7,
            schedule_seit_count_uploaded=7,
            schedule_athena_count_uploaded=7,
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            athena_avc=250,
            athena_ppa=3.324,
            validation_status="PENDING",
            created_by="tester",
        )
        self.db.add(run)
        self.db.commit()
        self.db.refresh(run)

        template_bytes = build_rewa_template_bytes(with_formulas=True)
        captured_kwargs = {}

        def fake_recalc(input_path, output_path, **kwargs):
            captured_kwargs.update(kwargs)
            build_recalculated_workbook(Path(output_path), (101.25, 202.5, None))
            return Path(output_path)

        cfg = {
            "schedule_generators": ["SPRNG", "SEIT", "ATHENA"],
            "workbook": {
                "sheet_name": "Aggregated_Calculations",
                "start_row": 4,
                "rows_per_day": 96,
                "days": 7,
                "input_columns": {
                    "date": "A",
                    "time_block": "B",
                    "SPRNG": {"schedule": "C", "actual": "D", "avc": "E", "ppa": "F"},
                    "SEIT": {"schedule": "R", "actual": "S", "avc": "T", "ppa": "U"},
                    "ATHENA": {"schedule": "AG", "actual": "AH", "avc": "AI", "ppa": "AJ"},
                },
            },
        }

        with patch("services.dsm_verification_service.get_pss_config", return_value=cfg), patch(
            "services.dsm_verification_service.get_active_template_row",
            return_value=SimpleNamespace(id=1, version=1, template_binary=template_bytes),
        ), patch(
            "services.dsm_verification_service.validate_run_inputs",
            return_value={"ok": True, "missing": []},
        ), patch(
            "services.dsm_verification_service._parse_input_files_for_run",
            return_value=({"ATHENA": {}}, {"ATHENA": {}}, {}, []),
        ), patch(
            "services.dsm_verification_service.recalculate_excel_workbook",
            side_effect=fake_recalc,
        ):
            result = generate_run_workbook(self.db, run.id)

        self.assertEqual(result.status, "COMPLETED")
        self.assertEqual(result.generated_filename, "REWA_DSM_Verification(2026)_2026-08-03_to_2026-08-09.xlsx")
        self.assertEqual(captured_kwargs.get("required_final_charge_cells"), {"Final_Summary": ("D7", "E7", "F7")})
        self.assertTrue(captured_kwargs.get("allow_formula_fallback"))

    def test_rewa_meter_upload_parses_athena_mr19_mr20(self):
        content = (
            "Block,MR19,MR20\n"
            "1,1.25,2.75\n"
            "2,3,4\n"
        ).encode("utf-8")

        parsed, errors = parse_meter_uploads_for_generators(
            "030826.QC13.csv",
            content,
            {
                "ATHENA": {
                    "source_columns": ["MR19", "MR20"],
                    "multiplier": 4.0,
                }
            },
        )

        self.assertFalse(errors)
        self.assertEqual(parsed["ATHENA"][1], 16.0)
        self.assertEqual(parsed["ATHENA"][2], 28.0)

    def test_rewa_validation_requires_configured_athena_schedule_and_values(self):
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            from_date=date(2026, 8, 3),
            to_date=date(2026, 8, 3),
            revision_number=1,
            status="DRAFT",
            meter_count_expected=1,
            meter_count_uploaded=0,
            schedule_sprng_count_uploaded=0,
            schedule_seit_count_uploaded=0,
            schedule_athena_count_uploaded=0,
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            validation_status="PENDING",
            created_by="tester",
        )
        self.db.add(run)
        self.db.commit()

        with patch("services.dsm_verification_service.get_active_template_row", return_value=SimpleNamespace(id=1)):
            result = validate_run_inputs(
                self.db,
                run,
                {"schedule_generators": ["SPRNG", "SEIT", "ATHENA"]},
            )

        self.assertIn("ATHENA Schedule\n03-Aug-2026", result["missing"])
        self.assertIn("ATHENA AVC", result["missing"])
        self.assertIn("ATHENA PPA", result["missing"])

    def test_populate_workbook_fills_athena_configured_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Aggregated_Calculations"
        wb.create_sheet("Final_Summary")
        wb.create_sheet("Summary_WA")
        wb.create_sheet("Summary_WoA")
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            from_date=date(2026, 8, 3),
            to_date=date(2026, 8, 3),
            revision_number=1,
            status="DRAFT",
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            athena_avc=250,
            athena_ppa=3.5,
            validation_status="PENDING",
            created_by="tester",
        )

        output = _populate_workbook(
            wb,
            cfg={
                "generator_defaults": {},
                "schedule_generators": ["SPRNG", "SEIT", "ATHENA"],
                "workbook": {
                    "sheet_name": "Aggregated_Calculations",
                    "start_row": 4,
                    "rows_per_day": 96,
                    "input_columns": {
                        "date": "A",
                        "time_block": "B",
                        "ATHENA": {"schedule": "AG", "actual": "AH", "avc": "AI", "ppa": "AJ"},
                    },
                },
            },
            run=run,
            schedules_by_generator={"ATHENA": {date(2026, 8, 3): {1: 11.0}}},
            actuals_by_generator={"ATHENA": {date(2026, 8, 3): {1: 12.0}}},
        )
        populated = load_workbook(BytesIO(output), data_only=False)
        sheet = populated["Aggregated_Calculations"]

        self.assertEqual(sheet["AG4"].value, 11.0)
        self.assertEqual(sheet["AH4"].value, 12.0)
        self.assertEqual(sheet["AI4"].value, 250.0)
        self.assertEqual(sheet["AJ4"].value, 3.5)
        self.assertEqual(
            sheet["Y4"].value,
            "=IF(R4<S4,IF((-R4+S4)>(0.35*T4),(0.1*E4),(S4-R4-W4-X4)),0)",
        )
        self.assertEqual(
            sheet["AN4"].value,
            "=IF(AG4<AH4,IF((-AG4+AH4)>(0.35*AI4),(0.1*T4),(AH4-AG4-AL4-AM4)),0)",
        )
        self.assertEqual(sheet["DB4"].value, "=SUM(CY4:CZ4:DA4)")
        self.assertTrue(str(sheet["AK4"].value).startswith("="))
        self.assertTrue(str(sheet["AZ4"].value).startswith("="))

        self.assertEqual(populated["Final_Summary"]["F6"].value, 250.0)
        self.assertEqual(populated["Summary_WA"]["F7"].value, "=Aggregated_Calculations!AJ4")
        self.assertEqual(
            populated["Final_Summary"]["F7"].value,
            "=(Summary_WA!F12+Summary_WA!F13)-(Summary_WA!F10+Summary_WA!F11)",
        )
        self.assertEqual(populated["Summary_WoA"]["G5"].value, "SI Units")
        self.assertEqual(populated["Summary_WoA"]["G6"].value, "MW")
        self.assertEqual(populated["Summary_WoA"]["G7"].value, "INR/kWh")
        self.assertEqual(populated["Summary_WoA"]["G8"].value, "INR")
        self.assertEqual(populated["Summary_WoA"]["G15"].value, "%")

    def test_populate_workbook_rounds_rewa_extended_athena_ppa_to_office_precision(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Aggregated_Calculations"
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            from_date=date(2026, 8, 3),
            to_date=date(2026, 8, 3),
            revision_number=1,
            status="DRAFT",
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            athena_avc=250,
            athena_ppa=3.324,
            validation_status="PENDING",
            created_by="tester",
        )

        output = _populate_workbook(
            wb,
            cfg={
                "schedule_generators": ["SPRNG", "SEIT", "ATHENA"],
                "workbook": {
                    "sheet_name": "Aggregated_Calculations",
                    "start_row": 4,
                    "rows_per_day": 96,
                    "input_columns": {
                        "date": "A",
                        "time_block": "B",
                        "ATHENA": {"schedule": "AG", "actual": "AH", "avc": "AI", "ppa": "AJ"},
                    },
                },
            },
            run=run,
            schedules_by_generator={"ATHENA": {date(2026, 8, 3): {1: 11.0}}},
            actuals_by_generator={"ATHENA": {date(2026, 8, 3): {1: 12.0}}},
        )
        populated = load_workbook(BytesIO(output), data_only=False)

        self.assertEqual(populated["Aggregated_Calculations"]["AJ4"].value, 3.32)

    def test_populate_workbook_detects_athena_columns_from_uploaded_template(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Aggregated_Calculations"
        ws["K1"] = "ATHENA"
        ws["K3"] = "Schedule"
        ws["L3"] = "Actual"
        ws["M3"] = "AVC"
        ws["N3"] = "PPA"
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            from_date=date(2026, 8, 3),
            to_date=date(2026, 8, 3),
            revision_number=1,
            status="DRAFT",
            athena_avc=250,
            athena_ppa=3.5,
            validation_status="PENDING",
            created_by="tester",
        )

        output = _populate_workbook(
            wb,
            cfg={
                "schedule_generators": ["ATHENA"],
                "workbook": {
                    "sheet_name": "Aggregated_Calculations",
                    "start_row": 4,
                    "rows_per_day": 96,
                    "input_columns": {
                        "date": "A",
                        "time_block": "B",
                        "ATHENA": {"schedule": "AG", "actual": "AH", "avc": "AI", "ppa": "AJ"},
                    },
                },
            },
            run=run,
            schedules_by_generator={"ATHENA": {date(2026, 8, 3): {1: 21.0}}},
            actuals_by_generator={"ATHENA": {date(2026, 8, 3): {1: 22.0}}},
        )
        populated = load_workbook(BytesIO(output), data_only=False)
        sheet = populated["Aggregated_Calculations"]

        self.assertEqual(sheet["K4"].value, 21.0)
        self.assertEqual(sheet["L4"].value, 22.0)
        self.assertEqual(sheet["M4"].value, 250.0)
        self.assertEqual(sheet["N4"].value, 3.5)
        self.assertIsNone(sheet["AG4"].value)

    def test_populate_workbook_keeps_2026_rewa_template_formulas_separate(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Aggregated_Calculations"
        wb.create_sheet("Final_Summary")
        wb.create_sheet("Summary_WA")
        wb.create_sheet("Summary_WoA")
        ws["C2"] = "Sprng"
        ws["P2"] = "Seit"
        ws["AC2"] = "Athena"
        ws["AP2"] = "Aggregated"
        ws["C3"] = "Schedule(MW)"
        ws["D3"] = "Actual(MW)"
        ws["E3"] = "AvC(MW)"
        ws["F3"] = "PPA"
        ws["P3"] = "Schedule(MW)"
        ws["Q3"] = "Actual(MW)"
        ws["R3"] = "AvC(MW)"
        ws["S3"] = "PPA"
        ws["AC3"] = "Schedule(MW)"
        ws["AD3"] = "Actual(MW)"
        ws["AE3"] = "AvC(MW)"
        ws["AF3"] = "PPA"
        ws["AP3"] = "Schedule(MW)"
        ws["AQ3"] = "Actual(MW)"
        ws["AR3"] = "AvC(MW)"
        ws["AS3"] = "PPA"
        ws["G4"] = "=IFERROR(ABS(D4-C4)*250,0)"
        ws["H4"] = "=IF(C4<D4,IF((-C4+D4)<(0.05*E4),(D4-C4),(0.05*E4)),0)"
        ws["T4"] = "=IFERROR(ABS(Q4-P4)*250,0)"
        ws["U4"] = "=IF(P4<Q4,IF((-P4+Q4)<(0.05*R4),(Q4-P4),(0.05*R4)),0)"
        ws["AG4"] = "=IFERROR(ABS(AD4-AC4)*250,0)"
        ws["AH4"] = "=IF(AC4<AD4,IF((-AC4+AD4)<(0.05*AE4),(AD4-AC4),(0.05*AE4)),0)"
        ws["AT4"] = "=IFERROR(ABS(AQ4-AP4)*250,0)"
        ws["AV4"] = "=IF(AP4<AQ4,IF((-AP4+AQ4)>(0.1*AR4),(0.05*AR4),(AQ4-AP4-AU4)),0)"
        wb["Summary_WA"]["G7"] = "=Aggregated_Calculations!AS4"
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            regulation="2026",
            from_date=date(2026, 8, 10),
            to_date=date(2026, 8, 10),
            revision_number=1,
            status="DRAFT",
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            athena_avc=250,
            athena_ppa=3.324,
            validation_status="PENDING",
            created_by="tester",
        )

        output = _populate_workbook(
            wb,
            cfg={
                "schedule_generators": ["SPRNG", "SEIT", "ATHENA"],
                "workbook": {
                    "sheet_name": "Aggregated_Calculations",
                    "start_row": 4,
                    "rows_per_day": 96,
                    "input_columns": {
                        "date": "A",
                        "time_block": "B",
                    },
                },
            },
            run=run,
            schedules_by_generator={
                "SPRNG": {date(2026, 8, 10): {1: 1.0}},
                "SEIT": {date(2026, 8, 10): {1: 2.0}},
                "ATHENA": {date(2026, 8, 10): {1: 3.0}},
            },
            actuals_by_generator={
                "SPRNG": {date(2026, 8, 10): {1: 1.5}},
                "SEIT": {date(2026, 8, 10): {1: 2.5}},
                "ATHENA": {date(2026, 8, 10): {1: 3.5}},
            },
        )
        populated = load_workbook(BytesIO(output), data_only=False)
        sheet = populated["Aggregated_Calculations"]

        self.assertEqual(sheet["P4"].value, 2.0)
        self.assertEqual(sheet["AC4"].value, 3.0)
        self.assertEqual(sheet["AP4"].value, "=SUM(C4,P4,AC4)")
        self.assertEqual(sheet["AQ4"].value, "=SUM(D4,Q4,AD4)")
        self.assertEqual(sheet["AS4"].value, "=AVERAGE(F4,S4,AF4)")
        self.assertEqual(sheet["H4"].value, "=IF(C4<D4,IF((-C4+D4)<(0.05*E4),(D4-C4),(0.05*E4)),0)")
        self.assertEqual(sheet["T5"].value, "=IFERROR(ABS(Q5-P5)*250,0)")
        self.assertEqual(sheet["U5"].value, "=IF(P5<Q5,IF((-P5+Q5)<(0.05*R5),(Q5-P5),(0.05*R5)),0)")
        self.assertEqual(sheet["AG5"].value, "=IFERROR(ABS(AD5-AC5)*250,0)")
        self.assertEqual(sheet["AH5"].value, "=IF(AC5<AD5,IF((-AC5+AD5)<(0.05*AE5),(AD5-AC5),(0.05*AE5)),0)")
        self.assertEqual(sheet["AV5"].value, "=IF(AP5<AQ5,IF((-AP5+AQ5)>(0.1*AR5),(0.05*AR5),(AQ5-AP5-AU5)),0)")
        self.assertEqual(
            populated["Final_Summary"]["D7"].value,
            "=(Summary_WA!D12+Summary_WA!D13)-(Summary_WA!D10+Summary_WA!D11)",
        )
        self.assertEqual(
            populated["Final_Summary"]["F7"].value,
            "=(Summary_WA!F12+Summary_WA!F13)-(Summary_WA!F10+Summary_WA!F11)",
        )
        self.assertEqual(
            populated["Summary_WA"]["D10"].value,
            "=D8-D9",
        )
        self.assertEqual(
            populated["Summary_WoA"]["D13"].value,
            '=ABS(SUMIFS(Aggregated_Calculations!BV:BV,Aggregated_Calculations!A:A,">="&Summary_WoA!E4,Aggregated_Calculations!A:A,"<="&Summary_WoA!F4))',
        )
        self.assertEqual(populated["Summary_WA"]["G7"].value, "=Aggregated_Calculations!AS4")

    def test_populate_workbook_fills_2026_rewa_freq_tras_date_range(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Aggregated_Calculations"
        wb.create_sheet("Final_Summary")
        wb.create_sheet("Summary_WA")
        wb.create_sheet("Summary_WoA")
        wb.create_sheet("Freq&TRAS")
        run = DsmVerificationRun(
            pss_code="REWA",
            pss_name="Rewa",
            regulation="2026",
            from_date=date(2026, 8, 17),
            to_date=date(2026, 8, 23),
            revision_number=1,
            status="DRAFT",
            sprng_avc=250,
            sprng_ppa=3.324,
            seit_avc=250,
            seit_ppa=3.279,
            athena_avc=250,
            athena_ppa=3.324,
            validation_status="PENDING",
            created_by="tester",
        )

        output = _populate_workbook(
            wb,
            cfg={
                "schedule_generators": ["SPRNG", "SEIT", "ATHENA"],
                "workbook": {
                    "sheet_name": "Aggregated_Calculations",
                    "start_row": 4,
                    "rows_per_day": 96,
                    "input_columns": {
                        "date": "A",
                        "time_block": "B",
                        "SPRNG": {"schedule": "C", "actual": "D", "avc": "E", "ppa": "F"},
                        "SEIT": {"schedule": "P", "actual": "Q", "avc": "R", "ppa": "S"},
                        "ATHENA": {"schedule": "AC", "actual": "AD", "avc": "AE", "ppa": "AF"},
                    },
                },
            },
            run=run,
            schedules_by_generator={},
            actuals_by_generator={},
        )
        populated = load_workbook(BytesIO(output), data_only=False)
        freq_tras = populated["Freq&TRAS"]

        self.assertEqual(freq_tras["A4"].value.date(), date(2026, 8, 17))
        self.assertEqual(freq_tras["B4"].value, "00:00-00:15")
        self.assertEqual(freq_tras["A99"].value.date(), date(2026, 8, 17))
        self.assertEqual(freq_tras["B99"].value, "23:45-23:59")
        self.assertEqual(freq_tras["A100"].value.date(), date(2026, 8, 18))
        self.assertEqual(freq_tras["B100"].value, "00:00-00:15")
        self.assertEqual(freq_tras["A675"].value.date(), date(2026, 8, 23))
        self.assertEqual(freq_tras["B675"].value, "23:45-23:59")


if __name__ == "__main__":
    unittest.main()
