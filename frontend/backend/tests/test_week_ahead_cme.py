import csv
import io
import os
import sys
import unittest
from datetime import date


BACKEND_DIR = os.path.dirname(os.path.dirname(__file__))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

import main


class CmeWeekAheadTests(unittest.TestCase):
    def test_cme_week_ahead_forecast_schedule_and_avc_are_normalized(self):
        values = [
            {"date": "2026-07-23", "block": 1, "declared_forecast": 1.25, "inter_avc": 0, "schedule": 9},
            {"date": "2026-07-22", "block": 1, "declared_forecast": 0, "inter_avc": 5, "schedule": 0},
            {"date": "2026-07-22", "block": 2, "declared_forecast": 0.5, "inter_avc": 0, "schedule": 7},
        ]

        normalized = main._week_ahead_normalize_cme_values(values, date(2026, 7, 21))

        self.assertEqual(normalized[0]["date"], "2026-07-22")
        self.assertEqual(normalized[0]["block"], 1)
        self.assertEqual(normalized[0]["declared_forecast"], 0)
        self.assertEqual(normalized[0]["inter_avc"], 0)
        self.assertEqual(normalized[0]["schedule"], 0)
        self.assertEqual(normalized[1]["date"], "2026-07-22")
        self.assertEqual(normalized[1]["block"], 2)
        self.assertEqual(normalized[1]["declared_forecast"], 0.5)
        self.assertEqual(normalized[1]["inter_avc"], 5)
        self.assertEqual(normalized[1]["schedule"], 0.5)

    def test_cme_week_ahead_csv_fills_declared_forecast_intra_avc_and_schedule(self):
        values = main._week_ahead_normalize_cme_values(
            [0, 0.75],
            date(2026, 7, 21),
        )
        template = (
            "Schedule Template for MH_VEDANJAY and revision WA,,,\n"
            ",Scheduling entity,MH_VEDANJAY,\n"
            ",Date,2026-07-21,\n"
            ",Revision No,WA,\n"
            "Block,Declared Forecast,Intra Avc,Schedule\n"
            "1,,,\n"
            "2,,,\n"
            "3,,,\n"
        ).encode("utf-8")

        output = main._week_ahead_fill_csv(template, values).decode("utf-8")

        self.assertIn("1,0,0,0", output)
        self.assertIn("2,0.75,5,0.75", output)


class OseplWeekAheadTests(unittest.TestCase):
    def test_osepl_week_ahead_inter_avc_is_zero_for_non_generation_blocks(self):
        values = [
            {"date": "2026-07-22", "block": 1, "declared_forecast": 0, "inter_avc": 20, "schedule": 0},
            {"date": "2026-07-22", "block": 26, "declared_forecast": 0.21, "inter_avc": 20, "schedule": 0.21},
        ]

        normalized = main._week_ahead_normalize_osepl_values(values)

        self.assertEqual(normalized[0]["inter_avc"], 0)
        self.assertEqual(normalized[1]["inter_avc"], 20)

    def test_osepl_week_ahead_csv_uses_zero_inter_avc_for_non_generation_blocks(self):
        values = main._week_ahead_normalize_osepl_values([
            {"date": "2026-07-22", "block": 1, "declared_forecast": 0, "inter_avc": 20, "schedule": 0},
            {"date": "2026-07-22", "block": 2, "declared_forecast": 0.5, "inter_avc": 20, "schedule": 0.5},
        ])
        template = (
            "Schedule Template for MH_VEDANJAY and revision WA,,,\n"
            ",Scheduling entity,MH_VEDANJAY,\n"
            ",Date,2026-07-22,\n"
            ",Revision No,WA,\n"
            "Block,Declared Forecast,Inter Avc,Schedule\n"
            "1,,,\n"
            "2,,,\n"
        ).encode("utf-8")

        output = main._week_ahead_fill_csv(template, values).decode("utf-8")

        self.assertIn("1,0,0,0", output)
        self.assertIn("2,0.5,20,0.5", output)


class ZetricWeekAheadTests(unittest.TestCase):
    def test_single_zetric_week_ahead_csv_declared_matches_two_schedule_columns(self):
        values = [{"block": idx + 1, "declared_forecast": 0.0887, "inter_avc": 17.2, "schedule": 0.0887} for idx in range(96)]
        template_lines = [
            "Schedule Template for MH_VEDANJAY and revision WA,,,,",
            ",Revision No,WA,,,",
            "POS Name,Chakur 132kV,Chakur 132kV,Chakur 132kV,Chakur 132kV",
            "Capacity,17.2,17.2,8.475,8.725",
            "Block,Declared Forecast,Intra Avc,Schedule,Schedule",
        ]
        template_lines.extend(f"{idx},,,," for idx in range(1, 97))
        template = ("\n".join(template_lines) + "\n").encode("utf-8")

        output = main._week_ahead_fill_csv(template, values, "ZETRIC").decode("utf-8")

        self.assertIn("1,0.08,17.2,0.04,0.04", output)


class JewliWeekAheadTests(unittest.TestCase):
    def test_jewli_forecast_megawatt_maps_to_declared_forecast(self):
        source = (
            "Timestamp (Asia/Kolkata),Timestamp (Asia/Kolkata),Forecast (MEGAWATT)\n"
            "2026-09-25 00:00,2026-09-25 00:15,94.984\n"
            "2026-09-25 00:15,2026-09-25 00:30,95.056\n"
        ).encode("utf-8")

        values = main._week_ahead_extract_values("enercast_Jewli_Weekahead.csv", source, plant_code="JEWLI")

        self.assertEqual(values[0]["date"], "2026-09-25")
        self.assertEqual(values[0]["block"], 1)
        self.assertEqual(values[0]["declared_forecast"], 94.984)
        self.assertEqual(values[1]["block"], 2)
        self.assertEqual(values[1]["declared_forecast"], 95.056)

    def test_jewli_single_csv_splits_declared_forecast_into_schedule_columns(self):
        values = [
            {"date": "2026-09-25", "block": 1, "declared_forecast": 94.984, "inter_avc": 0, "schedule": 94.984},
        ]
        template = (
            "Schedule Template for MH_VEDANJAY and revision WA,,,,,\n"
            "Capacity,194.4,194.4,7.2,93.6,93.6\n"
            "Block,Declared Forecast,Intra Avc,Schedule,Schedule,Schedule\n"
            "1,,,,,\n"
        ).encode("utf-8")

        output = main._week_ahead_fill_csv(template, values, "JEWLI").decode("utf-8")
        rows = list(csv.reader(io.StringIO(output)))
        data = rows[3]

        self.assertEqual(data[0], "1")
        self.assertEqual(float(data[1]), 94.98)
        self.assertEqual(float(data[2]), 100.8)
        split_total = sum(float(value) for value in data[3:6])
        self.assertAlmostEqual(split_total, 94.98, places=2)
        self.assertEqual(float(data[3]), 6.78)
        self.assertEqual(float(data[4]), 0)
        self.assertEqual(float(data[5]), 88.2)

    def test_jewli_week_ahead_uses_day_ahead_schedule_windows(self):
        values = [
            {"date": "2026-09-25", "block": 1, "declared_forecast": 100.8, "inter_avc": 0, "schedule": 100.8},
            {"date": "2026-09-25", "block": 26, "declared_forecast": 100.8, "inter_avc": 0, "schedule": 100.8},
            {"date": "2026-09-25", "block": 75, "declared_forecast": 100.8, "inter_avc": 0, "schedule": 100.8},
        ]
        template = (
            "Schedule Template for MH_VEDANJAY and revision WA,,,,,\n"
            "Capacity,194.4,194.4,7.2,93.6,93.6\n"
            "Block,Declared Forecast,Intra Avc,Schedule,Schedule,Schedule\n"
            "1,,,,,\n"
            "26,,,,,\n"
            "75,,,,,\n"
        ).encode("utf-8")

        output = main._week_ahead_fill_csv(template, values, "JEWLI").decode("utf-8")
        rows = list(csv.reader(io.StringIO(output)))

        self.assertEqual(rows[3][3:6], ["7.2", "0", "93.6"])
        self.assertEqual(rows[4][3:6], ["7.2", "0", "0"])
        self.assertEqual(rows[5][3:6], ["7.2", "93.6", "0"])


class TelanganaCombinedWeekAheadTests(unittest.TestCase):
    def test_zero_schedule_zeroes_avc_for_combined_section(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.cell(1, 1).value = "Block"
        sheet.cell(2, 1).value = 1
        sheet.cell(3, 1).value = 2

        wrote = main._week_ahead_fill_telangana_xlsx_sections(
            sheet,
            1,
            {"KASIPET": {"block_col": 1, "date_pairs": {"2026-09-14": (2, 3)}}},
            {"KASIPET": [
                {"date": "2026-09-14", "block": 1, "avc": 15, "schedule": 0},
                {"date": "2026-09-14", "block": 2, "avc": 15, "schedule": 0.01},
            ]},
        )

        self.assertTrue(wrote)
        self.assertEqual(sheet.cell(2, 2).value, 0)
        self.assertEqual(sheet.cell(2, 3).value, 0)
        self.assertEqual(sheet.cell(3, 2).value, 15)
        self.assertEqual(sheet.cell(3, 3).value, 0.01)


if __name__ == "__main__":
    unittest.main()
